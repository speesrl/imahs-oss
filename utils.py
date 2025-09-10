
import re
import os
import yaml
import time
import glob 
import magic
import joblib 
import socket
import logging
import hashlib
import threading
import traceback
from copy import copy
from tqdm import tqdm
from docx import Document
from PIL import Image
from typing import List
from PIL.ExifTags import TAGS
from functools import cache
from PyPDF2 import PdfReader
from openpyxl import load_workbook
from dataclasses import dataclass
from markdown_it import MarkdownIt

class FileInfo:
    def __init__(self, path, fast=False):
        assert os.path.isfile(path), f'{path} is not a valid file'
        self.fast = fast
        self.magic = magic.Magic(mime=True)
        self.path = path
        self.allowed = [attr for attr in dir(self) if (not isinstance(getattr(self.__class__, attr, None), property) and callable(getattr(self, attr)))]
    @property
    @cache
    def metadata(self):
        if self.fast:
            return self.sys()
        try:
            assert hasattr(self, self.type) and callable(getattr(self, self.type)), f'unknown file type: {self.type}'
            return {**self.sys(), **getattr(self, self.type)()}
        except Exception as e:
            return self.sys()
    @property
    @cache
    def type(self):
        for attr in self.allowed:
            if attr in  self.magic.from_file(self.path).split("/")[-1]:
                return attr
        return ''
    def pdf(self):
        with open(self.path, 'rb') as file:
            reader = PdfReader(file)
            metadata = reader.metadata
            return metadata

    def word(self):
        doc = Document(self.path)
        properties = doc.core_properties
        metadata = {
            'author': properties.author,
            'title': properties.title,
            'subject': properties.subject,
            'keywords': properties.keywords,
            'last_modified_by': properties.last_modified_by,
            'created': properties.created,
            'modified': properties.modified,
        }
        return metadata

    def excel(self):
        wb = load_workbook(self.path)
        properties = wb.properties
        metadata = {
            'title': properties.title,
            'subject': properties.subject,
            'author': properties.creator,
            'last_modified_by': properties.lastModifiedBy,
            'created': properties.created,
            'modified': properties.modified,
        }
        return metadata
    def image(self):
        image = Image.open(self.path)
        exif_data = image._lpopif()
        
        metadata = {}
        if exif_data:
            for tag, value in exif_data.items():
                tag_name = TAGS.get(tag, tag)
                metadata[tag_name] = value
        return metadata
    def sys(self):
        stats = os.stat(self.path)
        return {
            'size': stats.st_size, 
            'last_modified': time.ctime(stats.st_mtime),
            'created': time.ctime(stats.st_ctime),
            'last_accessed': time.ctime(stats.st_atime),
        }



class CharacterTextSplitter:
    def __init__(
        self,
        separator: str = "\n\n",
        chunk_size: int = 1000,
        chunk_overlap: int = 200
    ):
        self.separator = separator
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
    def split_text(self, text: str) -> List[str]:
        splits = text.split(self.separator)
        chunks: List[str] = []
        current = ""
        for part in splits:
            if len(part) > self.chunk_size:
                if current:
                    chunks.append(current)
                    current = ""
                start = 0
                while start < len(part):
                    end = start + self.chunk_size
                    sub = part[start:end]
                    chunks.append(sub)
                    start += self.chunk_size - self.chunk_overlap
                continue
            if not current:
                current = part
            else:
                candidate = current + self.separator + part
                if len(candidate) <= self.chunk_size:
                    current = candidate
                else:
                    chunks.append(current)
                    current = part
        if current:
            chunks.append(current)
        return chunks

class ContextAwareMarkdownSplitter:
    def separate_special_chars(self, text: str) -> str:
        text = re.sub(r"(\r?\n)", r" \1 ", text)
        text = re.sub(r'([^\w\s\n]+)', r' \1 ', text)
        text = re.sub(r"[ ]{2,}", " ", text)
        return text
    def remove_frontmatter(self, text: str) -> str:
        pattern = re.compile(r'^---\s*\n.*?\n---\s*\n', flags=re.DOTALL)
        return re.sub(pattern, '', text.strip())
    def __init__(
        self,
        chunk_size: int = 1000,
        chunk_overlap: int = 200,
    ):
        self._parser = MarkdownIt()
        self._base_splitter = CharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
        )
    def split_text(self, text: str) -> List[str]:
        text = self.remove_frontmatter(text)
        lines = text.splitlines()
        tokens = self._parser.parse(text)
        sections = []
        current_headers: List[str] = []
        buffer: List[str] = []
        def flush_buffer():
            if buffer:
                sections.append({
                    "headers": current_headers.copy(),
                    "content": "".join(buffer).strip()
                })
                buffer.clear()
        i = 0
        while i < len(tokens):
            token = tokens[i]
            if token.type == "heading_open":
                flush_buffer()
                level = int(token.tag[1])
                current_headers = current_headers[: level - 1]
                inline = tokens[i + 1]
                header_text = inline.content.strip()
                current_headers.append(header_text)
                i += 3
                continue
            if token.type == "table_open" and token.map:
                start, end = token.map
                table_text = "\n".join(lines[start:end]) + "\n"
                buffer.append(table_text)
                flush_buffer()
                while i < len(tokens) and tokens[i].type != "table_close":
                    i += 1
                i += 1
                continue
            if token.type == "fence" and token.content:
                buffer.append(token.content)
                buffer.append("\n")
                flush_buffer()
                i += 1
                continue
            if token.type == "paragraph_open":
                i += 1
                while i < len(tokens) and tokens[i].type != "paragraph_close":
                    if tokens[i].type == "inline":
                        buffer.append(tokens[i].content)
                    i += 1
                buffer.append("\n")
                i += 1
                continue
            if token.type == "inline" and token.content:
                buffer.append(tokens[i].content)
                buffer.append("\n")
                i += 1
                continue
            i += 1
        flush_buffer()
        chunks: List[str] = []
        pending_prefix = ""
        for sec in sections:
            prefix_lines = sec["headers"]
            prefix = "\n".join(prefix_lines) + ("\n\n" if prefix_lines else "")
            text_block = prefix + sec["content"]
            sub_chunks = self._base_splitter.split_text(text_block)
            for chunk in sub_chunks:
                chunk = self.separate_special_chars(chunk)
                body = chunk[len(prefix):].strip() if prefix else chunk.strip()
                if not body:
                    pending_prefix += chunk + ' \n '
                else:
                    merged = pending_prefix + chunk
                    chunks.append(merged)
                    pending_prefix = ""
        if pending_prefix:
            chunks.append(pending_prefix)
        return chunks
    async def async_split_text(self, text: str) -> List[str]:
        return self.split_text(text)
    
@dataclass
class Document:
    page_content: str 
    metadata: dict

class ContentLoader:
    def split(self, raw):
        pattern = r'^---\s*\n(.*?)\n---\s*\n(.*)$'
        match = re.match(pattern, raw, re.DOTALL)
        if match:
            frontmatter = match.group(1)
            content     = match.group(2)
            metadata    = yaml.safe_load(frontmatter)
            if 'source' not in metadata:
                metadata['source'] = None
            if 'checksum' not in metadata:
                metadata['checksum'] =  hashlib.blake2b(content.encode(), digest_size=16, usedforsecurity=False).hexdigest()
            return metadata, content
        else:
            return {'source': None}, raw
    def __init__(
        self,
        raw: str
    ):
        self.metadata, self.text = self.split(raw)
        super().__init__()
    def lazy_load(self):
        if len(self.text) > 1e5:
            splits = ContextAwareMarkdownSplitter(chunk_overlap=0, chunk_size=int(1e4)).split_text(self.text)
        else:
            splits = [ContextAwareMarkdownSplitter().remove_frontmatter(self.text)]
        for idx, section in enumerate(splits):
            yield Document(page_content=section + "  \n", metadata={'chunk_index': idx, **self.metadata})
    def load(self) -> list[Document]:
        return list(self.lazy_load())

class FileLoader:
    @property
    def metadata(self):
        self.load()
        return self.metadata_
    def __init__(
        self,
        file_path: str,
    ):
        self.file_path = file_path
        super().__init__()
    def lazy_load(self):
        with open(self.file_path, 'r') as f:
            raw = f.read()
            loader = ContentLoader(raw)
            self.metadata_ = copy(loader.metadata)
            assert os.path.splitext(os.path.split(self.file_path)[-1])[0] == self.metadata_['checksum'], f"checksums do not match {os.path.splitext(os.path.split(self.file_path)[-1])[0]} {self.metadata_['checksum']}"
            if self.metadata_.get('source', None) is None:
                self.metadata_['source'] = os.path.basename(self.file_path)
            yield from loader.lazy_load()
    def load(self) -> list[Document]:
        return list(self.lazy_load())

class FilesLoader:
    def __init__(self):
        self.FILE_LOADER_MAPPING = {
            ".md": (FileLoader, {}),
        }
        self.cache = {}
    def load_one(self, filepath: str) -> List[Document]:
        try:
            _, ext = os.path.splitext(filepath)
            if (
                os.path.getsize(filepath) == 0
                or ext not in self.FILE_LOADER_MAPPING
            ):
                return None, None
            clss, args = self.FILE_LOADER_MAPPING[ext]
            with open(filepath, 'rb', buffering=0) as f:
                loader = clss(filepath, **args)
                if loader:
                    splits = loader.load()
                    return splits, filepath, loader.metadata
        except Exception as e:
            logging.exception(e)
            return None, None, None
    def list(self, source_dir: str) -> List[Document]:
        res = {
            filepath
            for ext in self.FILE_LOADER_MAPPING
            for filepath in glob.glob(os.path.join(source_dir, f"**/*{ext}"), recursive=True)
        }
        res = [
            {
                filepath: FileInfo(filepath).sys()
            }
            for filepath in res
        ]
        return res 
    def load_all(self, source_dir: str, existing_files: List[str] = []) -> List[Document]:
        filepaths_ondisk = {
            os.path.splitext(os.path.split(filepath)[-1])[0]: filepath
            for ext in self.FILE_LOADER_MAPPING
            for filepath in glob.glob(os.path.join(source_dir, f"**/*{ext}"), recursive=True)
        }
        filepaths_toload = {
            filepaths_ondisk[checksum]
            for checksum in set([*filepaths_ondisk.keys()]) - set(existing_files)
        }
        documents = []
        n_splits = 0
        with joblib.parallel_config(backend='threading'):
            with tqdm(total=len(filepaths_toload), desc=f'Loading new documents from {source_dir} - filepaths_toload: {len(filepaths_toload)} filepaths_ondisk {len(filepaths_ondisk)}') as tkdm:
                for (splits, filepath, metadata) in joblib.Parallel(n_jobs=os.cpu_count()-2)(joblib.delayed(self.load_one)(file) for file in filepaths_toload):
                    if splits and metadata.get('checksum') not in existing_files:
                        documents += splits
                        n_splits += len(splits)
                    tkdm.set_postfix_str(f"n_splits: {n_splits}")
                    tkdm.update()
        return documents
        
class Singleton(type):
    _instances = {}
    _lock = threading.Lock()
    def __call__(cls, *args, **kwargs):
        if cls not in cls._instances:
            with cls._lock:
                cls._instances[cls] = super(Singleton, cls).__call__(*args, **kwargs)
        return cls._instances[cls]
    
def service2ip(name):
    try:
        return socket.gethostbyname(name)
    except socket.gaierror as e:
        pass

def moving_average(avg, x, alpha=.1):
    return alpha*x + (1-alpha)*avg


def where_am_i():
    stack = traceback.extract_stack()
    _, _, func, _ = stack[-2]
    return func